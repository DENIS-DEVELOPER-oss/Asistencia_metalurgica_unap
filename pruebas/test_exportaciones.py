"""
Excel, Word y PDF del reporte.

Los generadores reciben un diccionario, no modelos, así que se prueban con
`datos_de_prueba` y sin abrir MySQL. Se comprueba que el archivo salga bien
formado y que lleve de verdad los datos del curso: un reporte que se
descarga vacío o con las letras cambiadas no se nota hasta imprimirlo.
"""
import io
import zipfile

from django.test import SimpleTestCase
from openpyxl import load_workbook

from reportes.excel import generar_excel, generar_excel_general, generar_excel_sesion
from reportes.pdf import generar_pdf_reporte, generar_pdf_sesion
from reportes.word import generar_word, generar_word_sesion

from .datos_de_prueba import datos_de_una_sesion, reporte_de_grupo


def texto_del_docx(contenido):
    """Todo el texto de un .docx, sin depender de python-docx para leerlo."""
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        return z.read("word/document.xml").decode("utf-8")


def celdas_de_la_hoja(hoja):
    return [str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None]


class ExcelDelGrupo(SimpleTestCase):
    def setUp(self):
        self.contenido = generar_excel(reporte_de_grupo())
        self.libro = load_workbook(io.BytesIO(self.contenido))

    def test_es_un_xlsx_valido(self):
        self.assertTrue(self.contenido.startswith(b"PK"))

    def test_trae_hoja_de_resumen_y_de_detalle(self):
        self.assertEqual(len(self.libro.sheetnames), 2)

    def test_el_resumen_nombra_a_los_tres_alumnos(self):
        celdas = celdas_de_la_hoja(self.libro.worksheets[0])
        for nombre in ("QUISPE MAMANI JUAN CARLOS", "CONDORI APAZA MARIA LUZ",
                       "FLORES HUANCA ANA"):
            self.assertIn(nombre, celdas)

    def test_el_detalle_trae_una_columna_por_clase(self):
        celdas = celdas_de_la_hoja(self.libro.worksheets[1])
        for etiqueta in ("01/04", "02/04", "03/04", "04/04", "05/04"):
            self.assertIn(etiqueta, celdas)

    def test_el_detalle_trae_las_letras_de_asistencia(self):
        celdas = celdas_de_la_hoja(self.libro.worksheets[1])
        for letra in ("P", "F"):
            self.assertIn(letra, celdas)

    def test_lleva_la_cabecera_oficial_del_curso(self):
        celdas = " ".join(celdas_de_la_hoja(self.libro.worksheets[0]))
        self.assertIn("MET201", celdas)
        self.assertIn("ALVAREZ ROJAS ROSA MARIA", celdas)

    def test_sin_clases_dictadas_sigue_generando_el_archivo(self):
        vacio = reporte_de_grupo(total_sesiones=0)
        vacio["alumnos"] = []
        vacio["total_alumnos"] = 0
        self.assertTrue(generar_excel(vacio).startswith(b"PK"))


class ExcelConsolidado(SimpleTestCase):
    def test_una_hoja_general_mas_una_por_curso(self):
        reportes = [reporte_de_grupo(), reporte_de_grupo()]
        libro = load_workbook(io.BytesIO(generar_excel_general(reportes, "ALVAREZ ROSA")))
        self.assertEqual(len(libro.sheetnames), 3)

    def test_los_nombres_de_hoja_no_se_repiten(self):
        reportes = [reporte_de_grupo() for _ in range(4)]
        libro = load_workbook(io.BytesIO(generar_excel_general(reportes, "ALVAREZ ROSA")))
        self.assertEqual(len(libro.sheetnames), len(set(libro.sheetnames)))


class ElConsolidadoNoFirmaAQuienLoDescarga(SimpleTestCase):
    """
    El docente del reporte sale del grupo, nunca del usuario conectado.

    Una cuenta de administrador no tiene ficha de docente, así que
    `Usuario.nombre_completo` cae a su `username`: firmar el consolidado con
    ella lo dejaba a nombre de «admin».
    """

    def cabecera(self, contenido):
        hoja = load_workbook(io.BytesIO(contenido))["General"]
        filas = {}
        for fila in hoja.iter_rows(min_row=1, max_row=10):
            valores = [c.value for c in fila if c.value]
            if len(valores) >= 2 and str(valores[0]).endswith(":"):
                filas[valores[0]] = valores[1]
        return filas

    def test_un_solo_docente_se_nombra(self):
        cabecera = self.cabecera(generar_excel_general([reporte_de_grupo()], "admin"))
        self.assertEqual(cabecera["DOCENTE:"], "ALVAREZ ROJAS ROSA MARIA")

    def test_el_usuario_que_descarga_no_aparece_como_docente(self):
        cabecera = self.cabecera(generar_excel_general([reporte_de_grupo()], "admin"))
        self.assertNotIn("admin", str(cabecera.get("DOCENTE:")))
        self.assertEqual(cabecera["GENERADO POR:"], "admin")

    def test_con_varios_docentes_se_listan_todos(self):
        otro = reporte_de_grupo()
        otro["cabecera"] = dict(otro["cabecera"], docente="MAMANI QUISPE LUIS")
        cabecera = self.cabecera(generar_excel_general([reporte_de_grupo(), otro], "admin"))
        self.assertIn("DOCENTES:", cabecera)
        self.assertIn("ALVAREZ ROJAS ROSA MARIA", cabecera["DOCENTES:"])
        self.assertIn("MAMANI QUISPE LUIS", cabecera["DOCENTES:"])

    def test_con_muchos_docentes_se_resume(self):
        reportes = []
        for numero in range(6):
            r = reporte_de_grupo()
            r["cabecera"] = dict(r["cabecera"], docente=f"DOCENTE NUMERO {numero}")
            reportes.append(r)
        cabecera = self.cabecera(generar_excel_general(reportes, "admin"))
        self.assertIn("6 docentes", cabecera["DOCENTES:"])

    def test_la_tabla_lleva_una_columna_docente(self):
        hoja = load_workbook(io.BytesIO(generar_excel_general([reporte_de_grupo()], "admin")))["General"]
        titulos = []
        for fila in hoja.iter_rows(min_row=1, max_row=12):
            if any(c.value == "CURSO" for c in fila):
                titulos = [c.value for c in fila if c.value]
                break
        self.assertIn("DOCENTE", titulos)

    def test_sin_usuario_no_se_pinta_la_linea_de_generado_por(self):
        cabecera = self.cabecera(generar_excel_general([reporte_de_grupo()]))
        self.assertNotIn("GENERADO POR:", cabecera)


class ExcelDeUnaClase(SimpleTestCase):
    def test_trae_la_lista_y_la_observacion(self):
        libro = load_workbook(io.BytesIO(generar_excel_sesion(datos_de_una_sesion())))
        celdas = " ".join(celdas_de_la_hoja(libro.worksheets[0]))
        self.assertIn("QUISPE MAMANI JUAN CARLOS", celdas)
        self.assertIn("Avisó por teléfono", celdas)


class WordDelGrupo(SimpleTestCase):
    def setUp(self):
        self.xml = texto_del_docx(generar_word(reporte_de_grupo()))

    def test_lleva_a_los_alumnos_y_al_curso(self):
        self.assertIn("QUISPE MAMANI JUAN CARLOS", self.xml)
        self.assertIn("MET201", self.xml)

    def test_lleva_las_fechas_de_clase(self):
        self.assertIn("01/04", self.xml)
        self.assertIn("05/04", self.xml)

    def test_lleva_el_tema_de_cada_clase(self):
        self.assertIn("Tema de la clase 1", self.xml)


class WordDeUnaClase(SimpleTestCase):
    def test_lleva_la_lista_del_dia(self):
        xml = texto_del_docx(generar_word_sesion(datos_de_una_sesion()))
        self.assertIn("CONDORI APAZA MARIA LUZ", xml)


class PdfDelGrupo(SimpleTestCase):
    def test_es_un_pdf_valido(self):
        contenido = generar_pdf_reporte(reporte_de_grupo())
        self.assertTrue(contenido.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", contenido[-1024:])

    def test_con_muchas_clases_reparte_las_fechas_en_bloques(self):
        # Por encima de 8 fechas el reporte pasa a apaisado y parte la tabla;
        # es donde más fácil se rompe la maquetación.
        corto = generar_pdf_reporte(reporte_de_grupo(total_sesiones=4))
        largo = generar_pdf_reporte(reporte_de_grupo(total_sesiones=30))
        self.assertTrue(largo.startswith(b"%PDF-"))
        self.assertGreater(len(largo), len(corto))

    def test_sin_clases_dictadas_sigue_generando_el_archivo(self):
        vacio = reporte_de_grupo(total_sesiones=0)
        vacio["alumnos"] = []
        vacio["total_alumnos"] = 0
        self.assertTrue(generar_pdf_reporte(vacio).startswith(b"%PDF-"))


class PdfDeUnaClase(SimpleTestCase):
    def test_es_un_pdf_valido(self):
        contenido = generar_pdf_sesion(datos_de_una_sesion())
        self.assertTrue(contenido.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", contenido[-1024:])


def texto_del_pdf(contenido):
    """Texto de las páginas: ReportLab las guarda en ASCII85 y comprimidas."""
    import base64
    import re
    import zlib

    partes = []
    # ReportLab no deja salto de línea antes de `endstream`.
    for flujo in re.findall(rb"stream\r?\n(.*?)endstream", contenido, re.S):
        try:
            partes.append(zlib.decompress(base64.a85decode(flujo.strip(), adobe=True)))
        except Exception:  # noqa: BLE001 - imágenes y fuentes no son texto
            try:
                partes.append(zlib.decompress(flujo))
            except zlib.error:
                pass
    return b"".join(partes).decode("latin-1")


class SinLeyendaNiVicerrectorado(SimpleTestCase):
    """Se quitaron de todos los reportes a pedido del usuario."""

    QUITADOS = (
        "Leyenda", "LEYENDA", "VICERRECTORADO", "Vicerrectorado",
        "riesgo de inhabilitación", "Semáforo:", "El porcentaje es el de clases",
    )

    def comprobar(self, texto):
        for quitado in self.QUITADOS:
            self.assertFalse(quitado in texto, f"sigue apareciendo «{quitado}»")

    def test_pdf(self):
        for contenido in (
            generar_pdf_reporte(reporte_de_grupo()),
            generar_pdf_sesion(datos_de_una_sesion()),
        ):
            texto = texto_del_pdf(contenido)
            # Que el texto se lea de verdad: si no, la prueba pasaría siempre.
            self.assertIn("REPORTE DE ASISTENCIA", texto)
            self.comprobar(texto)

    def test_word(self):
        for contenido in (
            generar_word(reporte_de_grupo()),
            generar_word_sesion(datos_de_una_sesion()),
        ):
            # El cuerpo y también el pie de página, donde iba el Vicerrectorado.
            with zipfile.ZipFile(io.BytesIO(contenido)) as z:
                partes = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/")
                          and n.endswith(".xml")]
            self.assertTrue(any("footer" in n for n in z.namelist()))
            self.comprobar(" ".join(partes))

    def test_excel(self):
        for contenido in (
            generar_excel(reporte_de_grupo()),
            generar_excel_sesion(datos_de_una_sesion()),
            generar_excel_general([reporte_de_grupo()], "ALVAREZ ROSA"),
        ):
            libro = load_workbook(io.BytesIO(contenido))
            for hoja in libro.worksheets:
                self.comprobar(" ".join(celdas_de_la_hoja(hoja)))


class ColumnasDeTotalesDelPdf(SimpleTestCase):
    """«PRES.», «FALT.» y «% ASIST.» se montaban: tenían 7 mm de los tiempos de P/T/F/J."""

    def tablas_del_reporte(self, reporte):
        from unittest import mock

        import reportes.pdf as pdf

        tablas = []
        original = pdf.Table

        def registrar(datos, *args, **kwargs):
            tablas.append((datos, kwargs.get("colWidths")))
            return original(datos, *args, **kwargs)

        with mock.patch.object(pdf, "Table", registrar):
            pdf.generar_pdf_reporte(reporte)
        return [(d, a) for d, a in tablas if d and d[0] and d[0][-1] == "% ASIST."]

    def test_un_ancho_por_columna_y_espacio_para_los_titulos(self):
        from reportlab.lib.units import mm

        tablas = self.tablas_del_reporte(reporte_de_grupo())
        self.assertTrue(tablas)
        for datos, anchos in tablas:
            self.assertEqual(len(anchos), len(datos[0]))
            pres, falt, asist = anchos[-3:]
            self.assertGreaterEqual(pres, 12 * mm)
            self.assertGreaterEqual(falt, 12 * mm)
            self.assertGreaterEqual(asist, 16 * mm)

    def test_con_muchas_fechas_la_tabla_cabe_en_la_hoja(self):
        from reportlab.lib.pagesizes import A4, landscape

        from reportes.pdf import MARGEN

        reporte = reporte_de_grupo()
        base = reporte["sesiones"][0]
        reporte["sesiones"] = [
            dict(base, id_sesion=1000 + i, etiqueta=f"{i + 1:02d}/05") for i in range(40)
        ]
        for datos, anchos in self.tablas_del_reporte(reporte):
            self.assertLessEqual(sum(anchos), landscape(A4)[0] - 2 * MARGEN + 0.01)
