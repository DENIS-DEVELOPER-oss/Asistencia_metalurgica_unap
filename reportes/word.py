"""
Generación del reporte de asistencia en Word (.docx) con python-docx.

Mismo contenido que el PDF: encabezado institucional, datos del curso,
matriz alumno x fecha con las letras P / T / F / J, totales, porcentaje y
espacio para la firma del docente.
"""
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1E, 0x3A, 0x8A)
GRIS = RGBColor(0x47, 0x55, 0x69)
ROJO = RGBColor(0xB9, 0x1C, 0x1C)

# Relleno de fondo por letra de estado.
FONDO_ESTADO = {
    "P": "DCFCE7",
    "F": "FEE2E2",
}
TEXTO_ESTADO = {
    "P": RGBColor(0x15, 0x80, 0x3D),
    "F": RGBColor(0xB9, 0x1C, 0x1C),
}


def _sombrear(celda, color_hex):
    """Aplica color de fondo a una celda (python-docx no lo expone directamente)."""
    relleno = OxmlElement("w:shd")
    relleno.set(qn("w:val"), "clear")
    relleno.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(relleno)


def _texto(celda, valor, *, negrita=False, tamano=8, color=None, centrado=True):
    celda.text = ""
    parrafo = celda.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
    parrafo.paragraph_format.space_before = Pt(1)
    parrafo.paragraph_format.space_after = Pt(1)
    corrida = parrafo.add_run(str(valor))
    corrida.bold = negrita
    corrida.font.size = Pt(tamano)
    if color is not None:
        corrida.font.color.rgb = color
    return corrida


def _parrafo(documento, texto, *, tamano=9, negrita=False, color=None,
             alineacion=WD_ALIGN_PARAGRAPH.LEFT, espacio_despues=2):
    parrafo = documento.add_paragraph()
    parrafo.alignment = alineacion
    parrafo.paragraph_format.space_after = Pt(espacio_despues)
    corrida = parrafo.add_run(texto)
    corrida.bold = negrita
    corrida.font.size = Pt(tamano)
    if color is not None:
        corrida.font.color.rgb = color
    return parrafo


def _configurar_pagina(documento, apaisado):
    seccion = documento.sections[0]
    if apaisado:
        ancho, alto = seccion.page_width, seccion.page_height
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width, seccion.page_height = alto, ancho
    for margen in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, margen, Cm(1.2))
    return seccion


def _pie_de_pagina(seccion, cabecera):
    parrafo = seccion.footer.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = parrafo.add_run(
        f"{cabecera['codigo_curso']} · {cabecera['grupo']} · "
        f"{cabecera['periodo']}"
    )
    corrida.font.size = Pt(7)
    corrida.font.color.rgb = GRIS


def _bloque_encabezado(documento, cabecera, fecha_emision=None):
    _parrafo(
        documento, cabecera["universidad"], tamano=13, negrita=True, color=AZUL,
        alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=0,
    )
    _parrafo(
        documento, "REPORTE DE ASISTENCIA", tamano=11, negrita=True,
        alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=0,
    )
    _parrafo(
        documento, f"ESTUDIANTES[{cabecera['periodo']}]", tamano=9, negrita=True,
        color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=8,
    )

    datos = [
        ("FACULTAD:", cabecera["facultad"]),
        ("ESCUELA PROFESIONAL:", cabecera["escuela"]),
        ("CURSO:", f"{cabecera['codigo_curso']} - {cabecera['curso']}"),
        ("DOCENTE:", cabecera["docente"]),
        (
            "SEMESTRE:",
            f"{cabecera['semestre_texto']}  -  CRÉDITOS: {cabecera['creditos']}"
            f"  -  GRUPO: {cabecera['grupo']}",
        ),
    ]
    if fecha_emision is not None:
        datos.append(("FECHA DE EMISIÓN:", fecha_emision.strftime("%d/%m/%Y %H:%M")))

    tabla = documento.add_table(rows=len(datos), cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    for fila, (etiqueta, valor) in zip(tabla.rows, datos, strict=True):
        fila.cells[0].width = Cm(4.6)
        _texto(fila.cells[0], etiqueta, negrita=True, tamano=9, centrado=False)
        _texto(fila.cells[1], valor, tamano=9, centrado=False)
    documento.add_paragraph().paragraph_format.space_after = Pt(4)


def _bloque_firma(documento, nombre_docente):
    documento.add_paragraph().paragraph_format.space_after = Pt(24)
    _parrafo(
        documento, "_" * 42, tamano=9,
        alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=0,
    )
    _parrafo(
        documento, nombre_docente, tamano=9, negrita=True,
        alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=0,
    )
    _parrafo(
        documento, "Firma del docente", tamano=8, color=GRIS,
        alineacion=WD_ALIGN_PARAGRAPH.CENTER,
    )


def generar_word(reporte):
    """Devuelve los bytes del archivo .docx del reporte consolidado."""
    cabecera = reporte["cabecera"]
    sesiones = reporte["sesiones"]

    documento = Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(9)

    # Con muchas fechas el reporte se imprime en horizontal.
    seccion = _configurar_pagina(documento, apaisado=len(sesiones) > 8)
    _pie_de_pagina(seccion, cabecera)
    _bloque_encabezado(documento, cabecera, reporte.get("fecha_emision"))

    encabezados = ["N°", "CÓDIGO", "APELLIDOS Y NOMBRES"]
    encabezados += [s["etiqueta"] for s in sesiones]
    encabezados += ["PRES.", "FALT.", "% ASIST."]

    tabla = documento.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    fila_titulos = tabla.rows[0]
    for celda, titulo in zip(fila_titulos.cells, encabezados, strict=True):
        _sombrear(celda, "1E3A8A")
        _texto(celda, titulo, negrita=True, tamano=8, color=RGBColor(0xFF, 0xFF, 0xFF))
    # Repite la cabecera si la tabla se parte entre páginas.
    fila_titulos._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for alumno in reporte["alumnos"]:
        celdas = tabla.add_row().cells
        _texto(celdas[0], alumno["numero"], tamano=8)
        _texto(celdas[1], alumno["codigo"], tamano=8)
        _texto(celdas[2], alumno["nombre"], tamano=8, centrado=False)

        columna = 3
        for sesion in sesiones:
            letra = alumno["marcas"].get(sesion["id_sesion"], "")
            _texto(
                celdas[columna], letra, negrita=bool(letra), tamano=8,
                color=TEXTO_ESTADO.get(letra),
            )
            if letra in FONDO_ESTADO:
                _sombrear(celdas[columna], FONDO_ESTADO[letra])
            columna += 1

        for valor in (alumno["presentes"], alumno["faltas"]):
            _texto(celdas[columna], valor, tamano=8)
            columna += 1

        porcentaje = alumno["porcentaje_asistencia"]
        _texto(
            celdas[columna],
            "-" if porcentaje is None else f"{porcentaje:.2f}%",
            negrita=alumno["en_riesgo"], tamano=8,
            color=ROJO if alumno["en_riesgo"] else None,
        )
        if alumno["en_riesgo"]:
            _sombrear(celdas[columna], "FEE2E2")

    documento.add_paragraph().paragraph_format.space_after = Pt(2)
    if sesiones:
        detalle = " · ".join(
            f"{s['etiqueta']}: {s['tema'] or 'Sin tema'}" for s in sesiones
        )
        _parrafo(documento, f"Sesiones: {detalle}", tamano=7, color=GRIS)

    _bloque_firma(documento, cabecera["docente"])

    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def generar_word_sesion(datos):
    """Documento .docx con la lista de asistencia de una sola clase."""
    cabecera = datos["cabecera"]
    sesion = datos["sesion"]
    resumen = datos["resumen"]

    documento = Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(9)

    seccion = _configurar_pagina(documento, apaisado=False)
    _pie_de_pagina(seccion, cabecera)
    _bloque_encabezado(documento, cabecera, datos.get("fecha_emision"))

    hora_fin = sesion["hora_fin"].strftime("%H:%M") if sesion["hora_fin"] else "—"
    _parrafo(
        documento,
        f"CLASE: {sesion['fecha']:%d/%m/%Y}   ·   "
        f"HORA: {sesion['hora_inicio']:%H:%M} a {hora_fin}   ·   "
        f"ESTADO: {sesion['estado']}",
        tamano=9, negrita=True,
    )
    _parrafo(documento, f"TEMA: {sesion['tema'] or 'Sin tema registrado'}", tamano=9)

    porcentaje = resumen["porcentaje_asistencia"]
    _parrafo(
        documento,
        f"Presentes: {resumen['presentes']}   ·   "
        f"Faltas: {resumen['faltas']}   ·   "
        f"Total: {resumen['total']}   ·   Asistencia: "
        f"{'—' if porcentaje is None else f'{porcentaje:.2f}%'}",
        tamano=9, color=GRIS, espacio_despues=8,
    )

    encabezados = ["N°", "CÓDIGO", "APELLIDOS Y NOMBRES", "ESTADO", "OBSERVACIÓN"]
    tabla = documento.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    fila_titulos = tabla.rows[0]
    for celda, titulo in zip(fila_titulos.cells, encabezados, strict=True):
        _sombrear(celda, "1E3A8A")
        _texto(celda, titulo, negrita=True, tamano=8, color=RGBColor(0xFF, 0xFF, 0xFF))
    fila_titulos._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    anchos = [Cm(1.1), Cm(2.1), Cm(7.6), Cm(2.4), Cm(5.2)]
    for alumno in datos["alumnos"]:
        celdas = tabla.add_row().cells
        _texto(celdas[0], alumno["numero"], tamano=8)
        _texto(celdas[1], alumno["codigo"], tamano=8)
        _texto(celdas[2], alumno["nombre"], tamano=8, centrado=False)
        _texto(
            celdas[3], alumno["estado"],
            negrita=True, tamano=8, color=TEXTO_ESTADO.get(alumno["letra"]),
        )
        if alumno["letra"] in FONDO_ESTADO:
            _sombrear(celdas[3], FONDO_ESTADO[alumno["letra"]])
        _texto(celdas[4], alumno["observacion"], tamano=8, centrado=False)

        for celda, ancho in zip(celdas, anchos, strict=True):
            celda.width = ancho

    _bloque_firma(documento, cabecera["docente"])

    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()

